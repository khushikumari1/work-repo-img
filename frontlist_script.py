# -*- coding: utf-8 -*-
"""
frontlist_script.py - Main orchestration script for the Frontlist alt-text pipeline.

EXECUTION_MODE in config.py selects the backend:
  "local" - Gemini is called directly on this machine (no AWS).
  "s3"    - upload to S3, Lambda runs Gemini, results are polled back.

Output layout (one folder per book):
  OUTPUTs/<book>/<chapter>/  <ArtLog>with_img.xlsx  +  images/
  ARCHIVE/<book>/<chapter>/  original PDF, DOCX, ArtLog (moved after success)

Usage:  python frontlist_script.py [INPUT_FOLDER]
"""

import os
import sys
import json
import re
import uuid
import time
import shutil
import zipfile
import logging
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed

import pandas as pd
from PIL import Image
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as Image_
from openpyxl.utils.exceptions import InvalidFileException
from docx import Document
import pythoncom
import win32com.client

# ─── Ensure repo root is importable ──────────────────────────────────────────
_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)

# ─── Config ───────────────────────────────────────────────────────────────────
from config import (
    EXECUTION_MODE,
    TEMP_IMAGES_DIR,
    JSON_OUTPUT_DIR,
    TEMPLATE_EXCEL,
    MAX_WORKERS,
    BASE_DIR,
    OUT_BASE_DIRECTORY,
    # AWS settings
    AWS_ACCESS_KEY_ID,
    AWS_SECRET_ACCESS_KEY,
    AWS_REGION,
    S3_BUCKET,
    S3_IMAGE_PREFIX,
    S3_TEXT_PREFIX,
    S3_OUTPUT_PREFIX,
    S3_POLL_MAX_ATTEMPTS,
    S3_POLL_WAIT_SECONDS,
)

# ─── Local Gemini processor (imported lazily so AWS-only runs don't need GCP) ─
_SRC_DIR = os.path.join(_HERE, "gemini-alt-text-generation", "src")
if _SRC_DIR not in sys.path:
    sys.path.insert(0, _SRC_DIR)

# ─── PDF extractor ────────────────────────────────────────────────────────────
from pdf_image_extractor import extract_images_from_pdf

# ─── Logging ──────────────────────────────────────────────────────────────────
logger = logging.getLogger(__name__)
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(threadName)s | %(message)s",
    datefmt="%H:%M:%S",
)

# ─── AWS S3 client (only built in "s3" mode) ──────────────────────────────────
if EXECUTION_MODE == "s3":
    import boto3
    from botocore.exceptions import ClientError
    s3_client = boto3.client(
        "s3",
        aws_access_key_id=AWS_ACCESS_KEY_ID,
        aws_secret_access_key=AWS_SECRET_ACCESS_KEY,
        region_name=AWS_REGION,
    )
else:
    s3_client = None
bucket_name = S3_BUCKET

# ─── Output column template ───────────────────────────────────────────────────
_DEFAULT_TEMPLATE_COLUMNS = [
    "Figure Number", "Short Alt Text", "Long Alt Text", "Chapter Name/Number",
    "Page number", "Short Alt-text Length", "Long Alt-text Length",
    "Caption", "Category", "Confidence Score", "Relevance Score",
    "Correctness Score", "Accuracy Score",
]

if not os.path.exists(TEMPLATE_EXCEL):
    logger.warning(
        f"Template Excel not found at {TEMPLATE_EXCEL} - "
        "creating default template with standard columns."
    )
    pd.DataFrame(columns=_DEFAULT_TEMPLATE_COLUMNS).to_excel(TEMPLATE_EXCEL, index=False)

template_file = pd.read_excel(TEMPLATE_EXCEL)

logger.info(f"[MODE] EXECUTION_MODE = '{EXECUTION_MODE}'")


# ═════════════════════════════════════════════════════════════════════════════
# UTILITY HELPERS
# ═════════════════════════════════════════════════════════════════════════════

def tlog(stage, job=None):
    if job and "image_path" in job:
        logger.info(f"{stage} | {os.path.basename(job['image_path'])}")
    else:
        logger.info(stage)


def load_workbook_safely(path):
    """Open an openpyxl workbook, scrubbing any corrupt defined names."""
    try:
        wb = load_workbook(path)
    except FileNotFoundError:
        logger.error(f"Excel file not found: {path}")
        return None
    except InvalidFileException:
        logger.error(f"Invalid or corrupted Excel file: {path}")
        return None
    except zipfile.BadZipFile:
        logger.error(f"Excel file corrupted (bad zip): {path}")
        return None
    except Exception as e:
        logger.error(f"Unexpected error opening Excel: {path} | {e}")
        return None

    bad_defs = []
    for name in list(wb.defined_names):
        defn = wb.defined_names[name]
        if defn.attr_text is None or "#N/A" in str(defn.attr_text):
            bad_defs.append(name)
    for name in bad_defs:
        del wb.defined_names[name]
    if bad_defs:
        logger.warning(f"Removed invalid defined names: {bad_defs}")

    return wb


def clear_folder(folder_path):
    """Delete all contents of a folder without removing the folder itself."""
    if not os.path.exists(folder_path):
        logger.warning(f"Folder not found for cleanup: {folder_path}")
        return False
    try:
        for item in os.listdir(folder_path):
            item_path = os.path.join(folder_path, item)
            if os.path.isfile(item_path) or os.path.islink(item_path):
                os.remove(item_path)
            elif os.path.isdir(item_path):
                shutil.rmtree(item_path)
        return True
    except Exception as e:
        logger.error(f"Failed to clear folder {folder_path}: {e}")
        return False


def _normalise_fig_key(key: str) -> str:
    return key.strip().lower().replace("-", ".")


def _build_normalised_lookup(d: dict) -> dict:
    return {_normalise_fig_key(k): v for k, v in d.items()}


def extract_number_number_from_path(path):
    path_ = os.path.basename(path)
    file_name_removed = path.replace(path_, "")
    start_index = file_name_removed.lower().rfind("ch")
    end_index = file_name_removed.lower().rfind("\\")

    if end_index == -1:
        end_index = file_name_removed.lower().rfind("/")

    if start_index != -1 and end_index > start_index:
        return file_name_removed[start_index:end_index]

    # Fallback: immediate parent folder name
    return os.path.basename(os.path.dirname(path))


def _book_and_chapter(root, base):
    """Work out (book, chapter) names for chapter folder `root` found while scanning `base`."""
    base, root = os.path.normpath(base), os.path.normpath(root)
    rel = os.path.relpath(root, base)
    parts = [] if rel == "." else rel.split(os.sep)
    if not parts:                       # `base` itself is the chapter folder
        return os.path.basename(os.path.dirname(base)), os.path.basename(base)
    if len(parts) == 1:                 # `base` is the book folder
        return os.path.basename(base), parts[0]
    return parts[0], os.path.join(*parts[1:])   # `base` holds several books


# ═════════════════════════════════════════════════════════════════════════════
# DOCX PROCESSING
# ═════════════════════════════════════════════════════════════════════════════

def accept_all_changes(docx_path):
    """Use Word COM to accept all tracked changes and save the file."""
    pythoncom.CoInitialize()
    try:
        try:
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = False
            word.DisplayAlerts = False
        except Exception as e:
            logger.error(f"Failed to start Word COM | {e}")
            return False
        try:
            doc = word.Documents.Open(docx_path)
            doc.AcceptAllRevisions()
            doc.Save()
            doc.Close()
            return True
        except Exception as e:
            logger.error(f"Failed processing DOCX: {docx_path} | {e}")
            return False
        finally:
            word.Quit()
    finally:
        pythoncom.CoUninitialize()


def extract_text_from_docx(file_path):
    document = Document(file_path)
    merged_figure_dict = {}
    state = {"current_fig": None}

    for paragraph in document.paragraphs:
        para_text = paragraph.text.strip()
        if not para_text:
            continue
        merged_figure_dict, state = extract_figure_details(
            para_text, merged_figure_dict, state
        )
    return merged_figure_dict


def extract_figure_details(text, figures, state):
    pattern = re.compile(
        r"(Figure|Fig\.?)\s*(\d+(?:[.-]\d+)*)", re.IGNORECASE
    )
    match = pattern.search(text)

    if match:
        fig_number = match.group(2)
        fig_id = f"Figure {fig_number}"
        figures[fig_id] = ""
        state["current_fig"] = fig_id
        remaining = text[match.end():].strip()
        if remaining:
            figures[fig_id] += remaining + " "
        return figures, state

    if state["current_fig"]:
        figures[state["current_fig"]] += text + " "

    return figures, state


# ═════════════════════════════════════════════════════════════════════════════
# S3 MODE - upload / poll / download / delete
# ═════════════════════════════════════════════════════════════════════════════

def upload_files_to_s3(img_path):
    """Upload a .jpg and its matching .txt caption file to S3."""
    try:
        img_path = os.path.normpath(img_path)

        if not os.path.exists(img_path):
            logger.error(f"Image file does not exist: {img_path}")
            raise FileNotFoundError(f"Image file not found: {img_path}")

        img_name = os.path.basename(img_path)
        base_name, ext = os.path.splitext(img_name)

        if ext.lower() != ".jpg":
            logger.warning(f"File is not a JPG: {img_name}")
            return

        txt_path = os.path.normpath(
            os.path.join(os.path.dirname(img_path), f"{base_name}.txt")
        )

        if not os.path.exists(txt_path):
            logger.warning(f"No matching TXT file for {img_name}: {txt_path}")
            return

        # Upload text first so it exists before Lambda reads it
        s3_client.upload_file(txt_path, bucket_name, f"{S3_TEXT_PREFIX}{os.path.basename(txt_path)}")
        # Upload image - this triggers the Lambda
        s3_client.upload_file(img_path, bucket_name, f"{S3_IMAGE_PREFIX}{img_name}")

    except Exception as e:
        logger.error(f"Upload process failed for {img_path}: {str(e)}")
        raise

    return img_path


def check_and_download_s3_file(s3_key, local_path,
                                max_attempts=S3_POLL_MAX_ATTEMPTS,
                                wait_seconds=S3_POLL_WAIT_SECONDS):
    """Poll S3 for s3_key and download it once it appears."""
    s3_key = s3_key.strip().replace("\\", "/")
    attempt = 1

    while attempt <= max_attempts:
        try:
            s3_client.head_object(Bucket=bucket_name, Key=s3_key)
            s3_client.download_file(bucket_name, s3_key, local_path)
            return True
        except ClientError as e:
            if e.response["Error"]["Code"] != "404":
                logger.error(f"S3 error for {s3_key}: {e}")
                return False
            if attempt == max_attempts:
                logger.error(f"JSON timeout: {os.path.basename(s3_key)}")
                return False
            time.sleep(wait_seconds)
            attempt += 1
        except Exception as e:
            logger.error(f"Unexpected S3 error for {s3_key}: {e}")
            return False


def upload_excel_to_s3(excel_path):
    """Upload an Excel file to the root of the S3 bucket."""
    try:
        s3_client.upload_file(excel_path, bucket_name, os.path.basename(excel_path))
        return True
    except Exception as e:
        logger.error(f"Failed to upload Excel to S3: {excel_path} | {e}")
        return False


def delete_job_from_s3(job):
    """Delete the txt, jpg, and json objects for one job from S3."""
    try:
        img_name  = os.path.basename(job["image_path"])
        base_name = os.path.splitext(img_name)[0]

        keys = [
            f"{S3_TEXT_PREFIX}{base_name}.txt",
            f"{S3_IMAGE_PREFIX}{img_name}",
            f"{S3_OUTPUT_PREFIX}{base_name}.json",
        ]
        response = s3_client.delete_objects(
            Bucket=bucket_name,
            Delete={"Objects": [{"Key": k} for k in keys]},
        )
        for err in response.get("Errors", []):
            logger.error(f"S3 delete failed: {err['Key']} | {err['Message']}")
        return job
    except Exception as e:
        logger.error(f"S3 delete exception: {os.path.basename(job['image_path'])} | {e}")
        return None


def upload_job(job):
    try:
        upload_files_to_s3(job["image_path"])
        return job
    except Exception as e:
        logger.error(f"Upload failed: {os.path.basename(job['image_path'])} | {e}")
        return None


def wait_and_download_job(job):
    img_name  = os.path.basename(job["image_path"])
    base_name = os.path.splitext(img_name)[0]
    json_name = f"{base_name}.json"

    local_json = os.path.join(JSON_OUTPUT_DIR, json_name)

    success = check_and_download_s3_file(
        f"{S3_OUTPUT_PREFIX}{json_name}",
        local_json,
    )

    if not success:
        logger.error(f"JSON failed: {img_name}")
        return None

    job["json_path"] = local_json
    return job


# ═════════════════════════════════════════════════════════════════════════════
# LOCAL MODE - direct Gemini call
# ═════════════════════════════════════════════════════════════════════════════

def _process_job_locally(job):
    """Worker for local mode: call Gemini via local_processor, attach json_path."""
    from local_processor import process_image_locally  # lazy import

    result = process_image_locally(job["image_path"], job["caption"])

    if result.get("Error"):
        logger.error(f"Gemini failed: {os.path.basename(job['image_path'])}")
        return None

    base_name = os.path.splitext(os.path.basename(job["image_path"]))[0]
    job["json_path"] = os.path.join(JSON_OUTPUT_DIR, f"{base_name}.json")
    return job


# ═════════════════════════════════════════════════════════════════════════════
# CORE PIPELINE - generate_IDTF_in_dict
# ═════════════════════════════════════════════════════════════════════════════

def generate_IDTF_in_dict(images, fig_id_with_cap, dict_columns, chapter_number):
    """
    Core IDTF generation loop.  Behaviour depends on EXECUTION_MODE:
    "local" - calls Gemini directly in parallel threads.
    "s3"    - uploads to S3, polls for Lambda results, downloads JSONs.
    """

    pipeline_start = time.time()
    jobs = []

    print(f"\n  [IDTF] Images received from PDF: {len(images)}")
    for fig_key_from_pdf, image, _ in images:
        print(f"  [IDTF]   key='{fig_key_from_pdf}' size={image.size}")
    print()

    # ── STEP 1: PREPARE IMAGES + CAPTIONS ────────────────────────────────────
    unique_figs = {}
    for fig_key, image, _ in images:
        base_key = fig_key.replace("-", ".")
        parts = [str(int(p)) if p.isdigit() else p for p in base_key.split(".")]
        base_key = ".".join(parts)
        if base_key not in unique_figs:
            unique_figs[base_key] = image
        else:
            print(f"  [IDTF] Skipped duplicate: '{base_key}'")

    print(f"\n  [FINAL FIGS]: {list(unique_figs.keys())}\n")

    cap_lookup = _build_normalised_lookup(fig_id_with_cap)
    os.makedirs(TEMP_IMAGES_DIR, exist_ok=True)

    for fig_id, image in unique_figs.items():
        captions_ = ""
        try:
            norm_key  = _normalise_fig_key(f"Figure {fig_id.strip()}")
            captions_ = cap_lookup.get(norm_key, "")
        except Exception:
            captions_ = ""

        timestamp       = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        temp_image_name = f"merged_{timestamp}_{uuid.uuid4().hex[:8]}.jpg"
        img_path        = os.path.join(TEMP_IMAGES_DIR, temp_image_name)
        txt_path        = img_path.replace(".jpg", ".txt")

        image.save(img_path)
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write(captions_)

        jobs.append({
            "image_path": img_path,
            "txt_path":   txt_path,
            "caption":    captions_,
            "fig_id":     str(fig_id),
            "chapter":    chapter_number,
        })

    logger.info(f"Prepared {len(jobs)} images for processing")

    # ── STEP 2: PROCESS IMAGES (mode-dependent) ───────────────────────────────
    completed_jobs = []

    if EXECUTION_MODE == "local":
        logger.info("[LOCAL] Starting parallel Gemini processing...")

        results = [None] * len(jobs)
        with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
            futures = {
                executor.submit(_process_job_locally, job): i
                for i, job in enumerate(jobs)
            }
            for future in as_completed(futures):
                idx = futures[future]
                result = future.result()
                if result:
                    results[idx] = result

        completed_jobs = [r for r in results if r is not None]
        logger.info(
            f"[LOCAL] Gemini processing complete | "
            f"{len(completed_jobs)}/{len(jobs)} successful"
        )

    else:
        logger.info("[S3] Starting parallel upload to S3...")

        uploaded_jobs = []
        results = [None] * len(jobs)
        with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
            futures = {executor.submit(upload_job, job): i for i, job in enumerate(jobs)}
            for future in as_completed(futures):
                idx = futures[future]
                result = future.result()
                if result:
                    results[idx] = result
        uploaded_jobs = [r for r in results if r is not None]
        logger.info(f"[S3] Upload complete | {len(uploaded_jobs)} successful")

        logger.info("[S3] Waiting for Lambda JSON outputs...")
        results = [None] * len(uploaded_jobs)
        with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
            futures = {
                executor.submit(wait_and_download_job, job): i
                for i, job in enumerate(uploaded_jobs)
            }
            for future in as_completed(futures):
                idx = futures[future]
                result = future.result()
                if result:
                    results[idx] = result
        completed_jobs = [r for r in results if r is not None]
        logger.info(
            f"[S3] JSON processing complete | "
            f"{len(completed_jobs)}/{len(uploaded_jobs)} successful"
        )

    # ── STEP 3: EXCEL + TEMPLATE FILLING (serial, same for both modes) ────────
    excel_rows = []
    current_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    for job in completed_jobs:
        with open(job["json_path"], "r", encoding="utf-8") as f:
            data = json.load(f)

        img_complexity    = data.get("Image_complexity", "")
        short_text        = data.get("short_alt_text", "")
        long_text         = data.get("long_alt_text", "")
        confidence_score  = data.get("Confidence_Score", "")
        relevance_score   = data.get("Relevance_Score", "")
        correctness_score = data.get("Correctness_Score", "")
        accuracy_score    = data.get("Accuracy_Score", "")

        if img_complexity.upper() == "SIMPLE":
            long_text = ""

        caption_value = job["caption"] if job["caption"] and job["caption"].strip() else "NA"

        excel_rows.append({
            "Date":              current_time,
            "ISBN Name":         "MyBookISBN",
            "Figure Number":         job["fig_id"],
            "No of Chapters":    job["chapter"],
            "Image Name":        os.path.basename(job["image_path"]),
            "Confidence Score":  confidence_score,
            "Relevance Score":   relevance_score,
            "Correctness Score": correctness_score,
            "Accuracy Score":    accuracy_score,
            "Image Path":        job["image_path"],
            "short alt text":    short_text,
            "long alt text":     long_text,
            "caption":           caption_value,
            "Image Complexity":  img_complexity,
        })

        for k in dict_columns.keys():
            if k == "Short Alt Text":
                dict_columns[k].append(short_text)
            elif k == "Long Alt Text":
                dict_columns[k].append(long_text)
            elif k == "Chapter Name/Number":
                dict_columns[k].append(
                    job["chapter"].replace("ch", "Chapter ").replace("CH", "Chapter ")
                )
            elif k == "Page number":
                dict_columns[k].append("")
            elif k == "Short Alt-text Length":
                dict_columns[k].append(len(short_text))
            elif k == "Long Alt-text Length":
                dict_columns[k].append(len(long_text))
            elif k == "Caption":
                dict_columns[k].append(caption_value)
            elif k == "figure_path":
                dict_columns[k].append(job["image_path"])
            elif k == "Figure Number":
                fig_id_val = job["fig_id"]
                dict_columns[k].append(
                    fig_id_val if fig_id_val.startswith("FX") else f"Fig. {fig_id_val}"
                )
            elif k == "Category":
                dict_columns[k].append(img_complexity)
            elif k.strip() == "Confidence Score":
                dict_columns[k].append(confidence_score)
            elif k.strip() == "Relevance Score":
                dict_columns[k].append(relevance_score)
            elif k.strip() == "Correctness Score":
                dict_columns[k].append(correctness_score)
            elif k.strip() == "Accuracy Score":
                dict_columns[k].append(accuracy_score)
            else:
                dict_columns[k].append("")

    # ── STEP 4: WRITE RUNNING LOG EXCEL ──────────────────────────────────────
    log_excel = os.path.join(BASE_DIR, "FRONTLIST-IDTF.xlsx")

    if excel_rows:
        new_df = pd.DataFrame(excel_rows)

        if os.path.exists(log_excel):
            existing_df = pd.read_excel(log_excel)
            for col in new_df.columns:
                if col not in existing_df.columns:
                    existing_df[col] = ""
            for col in existing_df.columns:
                if col not in new_df.columns:
                    new_df[col] = ""
            new_df = new_df[existing_df.columns]
            final_df = pd.concat([existing_df, new_df], ignore_index=True)
        else:
            final_df = new_df

        final_df.to_excel(log_excel, index=False)

        if EXECUTION_MODE == "s3":
            upload_excel_to_s3(log_excel)
            logger.info("[S3] Running log uploaded to S3")
        else:
            logger.info(f"[LOCAL] Running log updated: {log_excel}")

    # ── STEP 5: CLEANUP ───────────────────────────────────────────────────────
    if EXECUTION_MODE == "s3":
        logger.info("[S3] Starting S3 cleanup...")
        with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
            futures = [executor.submit(delete_job_from_s3, job) for job in completed_jobs]
            for future in as_completed(futures):
                future.result()
        logger.info("[S3] S3 cleanup complete")
    else:
        # Delete caption + JSON temp files only. The temp IMAGES must stay until
        # generate_IDTF has copied them to images/ and embedded the thumbnails;
        # generate_IDTF clears the whole temp folder afterwards.
        logger.info("[LOCAL] Cleaning up temp files...")
        for job in completed_jobs:
            for path in (job.get("txt_path"), job.get("json_path")):
                if path and os.path.exists(path):
                    try:
                        os.remove(path)
                    except Exception as e:
                        logger.warning(f"Could not remove temp file {path}: {e}")

    # ── SUMMARY ───────────────────────────────────────────────────────────────
    total_time = round(time.time() - pipeline_start, 2)
    logger.info(
        f"Pipeline complete | {len(completed_jobs)} success | "
        f"{len(jobs) - len(completed_jobs)} failed | {total_time}s"
    )

    return dict_columns


# ═════════════════════════════════════════════════════════════════════════════
# OUTPUT EXCEL (with embedded image thumbnails)
# ═════════════════════════════════════════════════════════════════════════════

def generate_with_img_IDTF(excel_path, df_, out_dir=None):
    figure_path_list = df_["figure_path"].to_list()
    df_ = df_.drop(columns=["figure_path"])

    out_dir = out_dir or os.path.dirname(excel_path)
    base = os.path.splitext(os.path.basename(excel_path))[0]
    without_img = os.path.join(out_dir, base + "without_img.xlsx")
    with_img    = os.path.join(out_dir, base + "with_img.xlsx")

    df_.to_excel(without_img, index=False)

    wb = load_workbook(without_img)
    ws = wb.active
    thumbnail_col = "E"

    for index, fig_path_ in enumerate(figure_path_list):
        if fig_path_ and os.path.exists(fig_path_):
            img = Image_(fig_path_)
            img.width, img.height = 100, 100
            ws.add_image(img, f"{thumbnail_col}{index + 2}")
            ws.row_dimensions[index + 2].height = img.height * 0.75

    wb.save(with_img)
    os.remove(without_img)
    logger.info(f"Output saved: {with_img}")


# ═════════════════════════════════════════════════════════════════════════════
# MAIN PIPELINE CONTROLLER - one chapter folder
# ═════════════════════════════════════════════════════════════════════════════

def generate_IDTF(pdf_path, docx_path, excel_path, out_dir=None):
    """
    Full pipeline for a single chapter folder.
    Results (Excel + images/) are written to out_dir.
    Returns the number of figures processed.
    """
    logger.info("Starting IDTF generation")
    out_dir = out_dir or os.path.dirname(excel_path)
    os.makedirs(out_dir, exist_ok=True)
    clear_folder(TEMP_IMAGES_DIR)   # start from a clean staging folder

    # ── STEP 1: EXCEL INTEGRITY CHECK ────────────────────────────────────────
    logger.info("Validating Excel file...")
    sheet_total = []

    try:
        sheet = pd.read_excel(excel_path, engine="openpyxl", skiprows=8)
        if "TOTALS" not in sheet.columns:
            raise ValueError("'TOTALS' column not found")
        for val in sheet["TOTALS"].tolist():
            if val is None:
                break
            s = str(val).strip()
            if not s or s.lower() == "nan":
                break
            sheet_total.append(s)

    except Exception as e:
        logger.warning(f"Pandas failed to read Excel | {e}")
        logger.info("Attempting safe openpyxl fallback...")

        wb = load_workbook_safely(excel_path)
        if wb is None:
            raise Exception("Excel integrity check failed")

        try:
            ws = wb.active
            column_a_values = [cell.value for cell in ws["A"]]
            if "TOTALS" not in column_a_values:
                raise ValueError("'TOTALS' not found in Column A")
            fig_index_found = column_a_values.index("TOTALS")
            for val in column_a_values[fig_index_found + 1:]:
                if val is None:
                    break
                s = str(val).strip()
                if not s or s.lower() == "nan":
                    break
                sheet_total.append(s)
        except Exception as inner_e:
            logger.error(f"Excel structure invalid | {inner_e}")
            raise Exception("Excel structure corrupted")

    logger.info(f"Valid figure IDs loaded: {sheet_total}")

    # ── STEP 2: DOCX PROCESSING ───────────────────────────────────────────────
    logger.info("Processing DOCX file...")
    accept_all_changes(docx_path)
    fig_id_with_caption = extract_text_from_docx(docx_path)

    # ── STEP 3: IMAGE EXTRACTION ──────────────────────────────────────────────
    logger.info("Extracting images from PDF...")
    images = extract_images_from_pdf(pdf_path)
    logger.info(f"Extracted {len(images)} images from PDF")

    # Warn when the ArtLog lists figures the PDF did not contain
    expected = {m.group(0) for s in sheet_total if (m := re.match(r"\d+\.\d+", s))}
    found = {re.sub(r"[A-Za-z]+$", "", k.replace("-", ".")) for k, _, _ in images}
    missing_figs = sorted(expected - found, key=lambda s: [int(x) for x in s.split(".")])
    if missing_figs:
        logger.warning(f"ArtLog lists figures not found in PDF: {missing_figs}")

    chapter_number = extract_number_number_from_path(pdf_path)

    # ── STEP 4: IDTF GENERATION ───────────────────────────────────────────────
    dict_columns = {col.strip(): [] for col in template_file.columns}
    dict_columns["figure_path"] = []

    dict_columns = generate_IDTF_in_dict(
        images, fig_id_with_caption, dict_columns, chapter_number
    )

    df_ = pd.DataFrame(dict_columns)

    if df_.shape[0] < len(images):
        logger.warning(
            f"{len(images) - df_.shape[0]} of {len(images)} figures failed at Gemini - "
            "check the output before relying on it"
        )

    # ── STEP 4.5: SAVE IMAGES TO OUTPUT FOLDER ────────────────────────────────
    output_images_folder = os.path.join(out_dir, "images")
    os.makedirs(output_images_folder, exist_ok=True)

    for fig_id, fig_path in zip(dict_columns["Figure Number"], dict_columns["figure_path"]):
        if fig_path and os.path.exists(fig_path):
            safe_name = fig_id.replace("/", "_").replace("\\", "_").replace(":", "_")
            shutil.copy2(fig_path, os.path.join(output_images_folder, f"{safe_name}.jpg"))

    logger.info(f"Saved images to: {output_images_folder}")

    # ── STEP 5: SAVE OUTPUT EXCEL ─────────────────────────────────────────────
    generate_with_img_IDTF(excel_path, df_, out_dir)

    # ── STEP 6: CLEAR TEMP STAGING FOLDER ────────────────────────────────────
    clear_folder(TEMP_IMAGES_DIR)

    return df_.shape[0]


# ═════════════════════════════════════════════════════════════════════════════
# DIRECTORY SCANNER - entry point
# ═════════════════════════════════════════════════════════════════════════════

def generate_IDTF_first_step(path_):
    """
    Process every folder that contains a PDF + DOCX + ArtLog Excel.
    Results go to OUTPUTs/<book>/<chapter>; originals are moved to ARCHIVE/<book>/<chapter>.
    Returns the list of book output folders that received results.
    """
    path_ = os.path.normpath(path_)
    logger.info(f"Scanning directory: {path_}")
    total = success = failed = 0
    out_books = []

    for root, folders, files in os.walk(path_):
        folders[:] = [d for d in folders if d.lower() != "images"]

        pdf_files    = [f for f in files if f.lower().endswith(".pdf")]
        docx_files   = [f for f in files if f.lower().endswith(".docx") and not f.startswith("~$")]
        excel_all    = [f for f in files if f.lower().endswith((".xlsx", ".xls")) and not f.startswith("~$")]
        input_excel  = [f for f in excel_all if "with_img" not in f.lower()]
        output_excel = [f for f in excel_all if "with_img" in f.lower()]

        if not (pdf_files or docx_files or input_excel):
            if output_excel:
                logger.info(f"Already processed (output exists, no source files): {root}")
            continue

        missing = [n for n, found in (("PDF", pdf_files), ("DOCX", docx_files),
                                      ("ArtLog Excel", input_excel)) if not found]
        if missing:
            logger.warning(f"Skipping folder (missing {', '.join(missing)}): {root}")
            continue

        book, chapter = _book_and_chapter(root, path_)
        out_dir = os.path.join(OUT_BASE_DIRECTORY, book, chapter)

        total += 1
        pdf_path   = os.path.join(root, pdf_files[0])
        docx_path  = os.path.join(root, docx_files[0])
        excel_path = os.path.join(root, input_excel[0])
        logger.info(f"Processing folder: {root}")

        try:
            img_count = generate_IDTF(pdf_path, docx_path, excel_path, out_dir)
        except Exception as e:
            failed += 1
            logger.error(f"Folder processing failed: {root} | {e}")
            continue

        if img_count > 0:
            success += 1
            book_dir = os.path.join(OUT_BASE_DIRECTORY, book)
            if book_dir not in out_books:
                out_books.append(book_dir)

            archive_dir = os.path.join(BASE_DIR, "ARCHIVE", book, chapter)
            os.makedirs(archive_dir, exist_ok=True)
            for p in (pdf_path, docx_path, excel_path):
                try:
                    dest = os.path.join(archive_dir, os.path.basename(p))
                    if os.path.exists(dest):
                        os.remove(dest)
                    shutil.move(p, dest)
                except Exception as e:
                    logger.warning(f"Could not archive {p}: {e}")
            logger.info(f"Results: {out_dir} | originals archived: {archive_dir}")
        else:
            failed += 1
            logger.warning(f"No figures processed in {root} - source files kept")

    logger.info(f"Scan complete | {success} success | {failed} failed | {total} total folders")
    return out_books


# ═════════════════════════════════════════════════════════════════════════════
# CLI
# ═════════════════════════════════════════════════════════════════════════════

if __name__ == "__main__":
    input_folder = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE_DIR, "INPUTs")

    if not os.path.isdir(input_folder):
        print(f"ERROR: Input folder does not exist: {input_folder}")
        sys.exit(1)

    book_dirs = generate_IDTF_first_step(input_folder)

    if book_dirs:
        target = book_dirs[0] if len(book_dirs) == 1 else OUT_BASE_DIRECTORY
        print(f"\nResults: {target}")
        if hasattr(os, "startfile"):          # Windows: open the folder for checking
            try:
                os.startfile(target)
            except Exception:
                pass